import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

VEHICLE_HEADERS = ["Ref", "Consignee", "Model", "Reg", "Location", "Inst", "Invoice", "V5", "Received", "Notes"]


def vehicle_row(vehicle):
    return [
        vehicle.ref,
        vehicle.consignee,
        vehicle.model,
        vehicle.reg,
        vehicle.location,
        "X" if vehicle.inst else "",
        "X" if vehicle.invoice else "",
        "X" if vehicle.v5 else "",
        vehicle.received,
        vehicle.notes,
    ]


def build_trip_xlsx(trip) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Yuk Listesi"

    ws.append([f"Booking Ref: {trip.booking_ref}"])
    ws.append(
        [f"Console: {trip.console}", f"Container: {trip.container}", f"ETS: {trip.ets}", f"ETA: {trip.eta}"]
    )
    ws.append([])
    ws.append(VEHICLE_HEADERS)
    for vehicle in trip.vehicles:
        ws.append(vehicle_row(vehicle))

    for i in range(1, len(VEHICLE_HEADERS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_trip_pdf(trip) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()

    elements = [Paragraph("Yuk Listesi", styles["Title"])]
    info = (
        f"Booking Ref: {trip.booking_ref} | Console: {trip.console} | "
        f"Container: {trip.container} | ETS: {trip.ets} | ETA: {trip.eta}"
    )
    elements.append(Paragraph(info, styles["Normal"]))
    elements.append(Spacer(1, 12))

    data = [VEHICLE_HEADERS] + [vehicle_row(v) for v in trip.vehicles]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buffer.getvalue()
